"""
Заявка перевозчику.

Госформата у заявки нет — это коммерческий документ. Поэтому здесь не XSD,
а обычный шаблон, и два выхода:

  • DOCX — распечатать, подписать, отправить перевозчику;
  • текст — кинуть в мессенджер, если перевозчик «свой» и бумага потом.

Зачем это в боте про накладные: заявка идёт ПЕРЕД накладной. Сначала заказали
машину, потом выписали ЭТрН. Если заявку сформировал бот, то к моменту
накладной он уже знает маршрут, груз, перевозчика и сроки — распознавать
остаётся только водителя и машину.
"""

from __future__ import annotations

from datetime import datetime

from .fields import BY_KEY, ZAYAVKA, registry_for
from .resolve import Carriage


def _fmt(v: object) -> str:
    if v is None or v == "":
        return "—"
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y %H:%M")
    # Ставка из парсера денег приходит float'ом (85000.0) — пишем её так же,
    # как целые: «85 000», а не «85000» рядом с «20 000 кг».
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (int, float)):
        return f"{v:,}".replace(",", " ")
    return str(v)


def as_text(c: Carriage) -> str:
    """Версия для мессенджера. Компактная, без воды."""
    g = c.get
    rate = _fmt(g("rate"))
    lines = [
        f"ЗАЯВКА № {_fmt(g('doc_number'))} от {_fmt(g('doc_date'))}",
        "",
        f"Заказчик: {_fmt(g('shipper_name'))}, ИНН {_fmt(g('shipper_inn'))}",
        f"Перевозчик: {_fmt(g('carrier_name'))}, ИНН {_fmt(g('carrier_inn'))}",
        "",
        f"Погрузка: {_fmt(g('pickup_address'))}",
        f"  {_fmt(g('pickup_datetime'))}, {_fmt(g('loading_type'))} загрузка",
        f"  контакт: {_fmt(g('pickup_contact'))}",
        f"Выгрузка: {_fmt(g('delivery_address'))}",
        f"  {_fmt(g('delivery_datetime'))}",
        f"  контакт: {_fmt(g('delivery_contact'))}",
        "",
        f"Груз: {_fmt(g('cargo_name'))}",
        f"  {_fmt(g('cargo_places'))} мест, {_fmt(g('cargo_package'))}, "
        f"{_fmt(g('cargo_mass_gross_kg'))} кг",
        f"Транспорт: {_fmt(g('req_vehicle_type'))}, "
        f"{_fmt(g('req_capacity_t'))} т, {_fmt(g('req_volume_m3'))} м³",
    ]
    if g("temperature_mode"):
        lines.append(f"Температурный режим: {_fmt(g('temperature_mode'))}")
    if g("special_instructions"):
        lines.append(f"Особые условия: {_fmt(g('special_instructions'))}")
    lines += [
        "",
        f"Ставка: {rate} ₽, {_fmt(g('rate_vat'))}",
        f"Оплата: {_fmt(g('payment_terms'))}",
        f"Простой: {_fmt(g('free_hours'))} ч включено, "
        f"далее {_fmt(g('demurrage'))} ₽/сутки",
    ]
    return "\n".join(lines)


# Стандартные условия. Правятся под свою практику в настройках организации,
# а не в коде — поэтому вынесены в отдельную константу.
DEFAULT_TERMS = [
    "Перевозчик подтверждает заявку подписью и печатью. "
    "Подписанная заявка имеет силу договора на данную перевозку.",
    "Перевозчик обязан подать исправное транспортное средство, "
    "пригодное для перевозки заявленного груза, к указанному времени.",
    "Замена транспортного средства или водителя без письменного "
    "согласования с Заказчиком не допускается.",
    "Перевозчик несёт ответственность за сохранность груза "
    "с момента приёма к перевозке и до момента выдачи грузополучателю.",
    "Оригиналы транспортной накладной и иных документов направляются "
    "Заказчику в течение 5 рабочих дней с даты выгрузки.",
]


def as_docx(c: Carriage, terms: list[str] | None = None) -> bytes:
    """DOCX на подпись. Строится программно — шаблон-файл не нужен."""
    import io

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    g = c.get
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(f"ЗАЯВКА № {_fmt(g('doc_number'))} на перевозку груза")
    run.bold = True
    run.font.size = Pt(14)

    sub = doc.add_paragraph(f"от {_fmt(g('doc_date'))}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    blocks: list[tuple[str, list[str]]] = [
        ("Стороны", ["shipper_name", "shipper_inn", "shipper_address", "shipper_phone",
                     "carrier_name", "carrier_inn", "carrier_address", "carrier_phone"]),
        ("Маршрут", ["pickup_address", "pickup_datetime", "loading_type", "pickup_contact",
                     "delivery_address", "delivery_datetime", "delivery_contact"]),
        ("Груз", ["cargo_name", "cargo_places", "cargo_package", "cargo_mass_gross_kg",
                  "cargo_volume_m3", "cargo_declared_value", "temperature_mode",
                  "special_instructions"]),
        ("Транспорт", ["req_vehicle_type", "req_capacity_t", "req_volume_m3"]),
        ("Оплата и простой", ["rate", "rate_vat", "payment_terms",
                              "free_hours", "demurrage"]),
    ]

    known = {f.key for f in registry_for(ZAYAVKA)}
    for title, keys in blocks:
        p = doc.add_paragraph()
        p.add_run(title).bold = True
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.columns[0].width = Cm(6)
        for key in keys:
            if key not in known or c.get(key) in (None, ""):
                continue
            row = table.add_row().cells
            row[0].text = BY_KEY[key].title
            row[1].text = _fmt(c.get(key))

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Условия").bold = True
    for i, t in enumerate(terms or DEFAULT_TERMS, 1):
        doc.add_paragraph(f"{i}. {t}")

    doc.add_paragraph()
    sig = doc.add_table(rows=2, cols=2)
    sig.cell(0, 0).text = "Заказчик"
    sig.cell(0, 1).text = "Перевозчик"
    sig.cell(1, 0).text = "\n____________________ / ______________ /\nМ.П."
    sig.cell(1, 1).text = "\n____________________ / ______________ /\nМ.П."

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
